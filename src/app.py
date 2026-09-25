import pandas as pd
import unicodedata
import dateparser
from datetime import datetime
from openpyxl.styles import Font, PatternFill
from dotenv import load_dotenv
from typesafe_sdk import TypeSafeClient, Choice, TypeSafeError
import phonenumbers
from pathlib import Path as pat
from tqdm import tqdm
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
import calendar

pd.set_option("display.max_columns", None)
pd.set_option("display.width", 200)

file_date_orders = {}

column_names = {
    "問い合わせID": "id",
    "受付日": "date",
    "氏名": "name",
    "会社名": "company",
    "電話番号": "phone",
    "金額": "amount",
    "問い合わせ内容": "content",

    "N° demande": "id",
    "Date de réception": "date",
    "Nom complet": "name",
    "Société": "company",
    "Téléphone": "phone",
    "Montant": "amount",
    "Message": "content",

    "Inquiry ID": "id",
    "Date Received": "date",
    "Full Name": "name",
    "Company": "company",
    "Phone": "phone",
    "Amount": "amount",
}

required_columns = ["id", "date", "name", "company", "phone", "amount", "content"]

offices_setting = pd.read_excel("settings.xlsx", sheet_name="Offices")
countries_setting = pd.read_excel("settings.xlsx", sheet_name="Countries")

office_regions = dict(zip(offices_setting["office"], offices_setting["country"]))
date_orders = dict(zip(countries_setting["country"], countries_setting["date_order"]))
default_currencies = dict(zip(countries_setting["country"], countries_setting["currency"]))
decimal_marks = dict(zip(countries_setting["country"], countries_setting["decimal_mark"]))

for office, country in office_regions.items():
    if country not in date_orders:
        raise SystemExit(f"settings.xlsx: office '{office}' uses country '{country}', " f"which is not in the Countries sheet.")

level_colors = {
    "Error": "F8CBAD",
    "Check": "FFE699",
    "Info": "D9D9D9",
    "AI": "BDD7EE"
}

JEV_THRESHOLD = 0.7

load_dotenv()
jev = TypeSafeClient()


def clean_text(value):
    if isinstance(value, str):
        return unicodedata.normalize("NFKC", value).strip()
    return value

def find_office(file_stem):
    for office in office_regions:
        keyword = office.lower().replace(" ", "")
        if keyword in file_stem.lower():
            return office
    return None

def prove_date_order(value):
    parts = value.split("/")
    if len(parts) != 3:
        return None
    if not parts[0].isdigit() or not parts[1].isdigit() or not parts[2].isdigit():
        return None
    if len(parts[2]) != 4:
        return None

    first = int(parts[0])
    second = int(parts[1])
    year = int (parts[2])

    dmy = False
    if 1 <= second <= 12:
        if 1 <= first <= calendar.monthrange(year, second)[1]:
            dmy = True

    mdy = False
    if 1 <= first <= 12:
        if 1 <= second <= calendar.monthrange(year, first)[1]:
            mdy = True

    if dmy and not mdy:
        return "DMY"
    if mdy and not dmy:
        return "MDY"

    return None

def find_file_date_orders(table):
    orders = {}
    for source_file, group in table.groupby("source_file"):
        found = set()
        for value in group["date"]:
            if isinstance(value, str):
                order = prove_date_order(value)
                if order is not None:
                    found.add(order)
        if len(found) == 1:
            orders[source_file] = found.pop()
        elif len(found) > 1:
            office = group["office"].iloc[0]
            issues.append({"office": office,
                        "id": f"(whole file: {source_file})",
                        "problem": "The file mixes two date orders. Its dates were read with the office format only"})
            orders[source_file] = None
        else:
            orders[source_file] = None

    return orders

def parse_date(row):
    value = row["date"]
    office = row["office"]
    row_id = row["id"]

    if isinstance(value, datetime):
        return value
    if not isinstance(value, str):
        issues.append({"office": office, "id": row_id,
                       "problem": "Date is missing"})
        return None

    if value.startswith("R"):
        year, month, day = value[1:].split(".")
        return datetime(2018 + int(year), int(month), int(day))

    proven_order = None
    if value[:4].isdigit():
        order = "YMD"
    else:
        proven_order = file_date_orders.get(row["source_file"])
        if proven_order is not None:
            order = proven_order
        else:
            order = date_orders[office_regions[office]]

    result = dateparser.parse(
        value,
        languages=["ja", "fr", "en"],
        settings={"DATE_ORDER": order},
    )

    if result is None:
        issues.append({"office": office, "id": row_id,
                       "problem": f"Could not read date '{value}'"})
        return None

    if str(result.year) not in value:
        issues.append({"office": office, "id": row_id,
                       "problem": f"Year missing or abbreviated in '{value}' (assumed {result.year})"})

    parts = value.split("/")
    if len(parts) == 3 and order in ["MDY", "DMY"]:
        first = int(parts[0])
        second = int(parts[1])
        if order == "MDY":
            expected = result.month
        else:
            expected = result.day
        if first != expected:
            issues.append({"office": office, "id": row_id,
                           "problem": f"Date '{value}' does not follow the {office} format ({order})"})

        if first <= 12 and second <= 12 and first != second:
            if order == "MDY":
                other_reading = datetime(result.year, second, first)
            else:
                other_reading = datetime(result.year, first, second)

            if other_reading <= datetime.now():
                if proven_order is None:
                    issues.append({"office": office, "id": row_id, "problem": f"Ambiguous date '{value}': read as {result:%Y-%m-%d} " f"({office} format), could also be {other_reading:%Y-%m-%d}"})
                else:
                    issues.append({"office": office, "id": row_id,
                                   "problem": f"Date '{value}' has two possible readings, but {proven_order} is "
                                              f"confirmed by the other dates in the same file "
                                              f"(read as {result:%Y-%m-%d})"})

    if result > datetime.now():
        issues.append({"office": office, "id": row_id,
                       "problem": f"Date '{value}' is in the future (read as {result:%Y-%m-%d})"})

    return result


def parse_currency(row):
    office = row["office"]
    amount = row["amount"]

    if pd.isna(amount):
        return None
    if not isinstance(amount, str):
        return default_currencies[office_regions[office]]

    if "¥" in amount or "円" in amount:
        return "JPY"
    elif "€" in amount:
        return "EUR"
    elif "$" in amount or "USD" in amount:
        return "USD"
    else:
        return default_currencies[office_regions[office]]

def ask_jev_value(office, big_value, small_value, content):
    state = (f"Office: {office}. A number in this message could mean either "
             f"{big_value} or {small_value}. Message: {content}")
    try:
        response = jev.system_one(
            state=state,
            questions={
                "value": Choice(
                    instructions="Which value does the customer mean?",
                    criteria={
                        "thousands": f"{big_value} (a large number)",
                        "decimal": f"{small_value} (a small number with decimals)",
                    },
                ),
            },
        )
    except TypeSafeError as error:
        tqdm.write(f"Jev error: {error}")
        return None, 0.0

    answer = response.answers["value"]
    return answer.choice, answer.confidence

def name_candidates(name):
    candidates = {}
    for i in range(1, len(name)):
        candidates[f"split_{i}"] = f"Family name: {name[:i]} / Given name: {name[i:]}"
    candidates["family_only"] = f"Family name only: {name} (the given name is missing)"
    return candidates

def ask_jev_name(name):
    candidates = name_candidates(name)
    state = f"A Japanese person's name was written without a space between family name and given name: {name}"
    try:
        response = jev.system_one(
            state=state,
            questions={
                "split": Choice(
                    instructions="How should this name be divided into family name and given name?",
                    criteria=candidates,
                ),
            },
        )
    except TypeSafeError as error:
        tqdm.write(f"Jev error: {error}")
        return None, 0.0

    answer = response.answers["split"]
    return answer.choice, answer.confidence


def parse_amount(row):
    office = row["office"]
    amount = row["amount"]
    row_id = row["id"]

    if pd.isna(amount):
        issues.append({"office": office, "id": row_id, "problem": "Amount is missing"})
        return None
    if not isinstance(amount, str):
        return float(amount)

    number_text = ""
    for c in amount:
        if c in "0123456789,.-":
            number_text = number_text + c

    if "," in number_text and "." in number_text:
        if number_text.rfind(",") > number_text.rfind("."):
            decimal = ","
        else:
            decimal = "."
    else:
        decimal = decimal_marks[office_regions[office]]

    parts = number_text.split(decimal)
    ambiguous = len(parts) == 2 and len(parts[1]) == 3

    if decimal == ",":
        number_text = number_text.replace(".", "")
        number_text = number_text.replace(",", ".")
    else:
        number_text = number_text.replace(",", "")


    try:
        result = float(number_text)
    except ValueError:
        issues.append({"office": office, "id": row_id, "problem": f"Could not read amount '{amount}'"})
        return None

    if result < 0:
        issues.append({"office": office, "id": row_id, "problem": f"Amount is negative ({amount}). Not accepted"})
        return None
    if ambiguous:
        big_value = int(round(result * 1000))
        choice, confidence = ask_jev_value(office,big_value, result, row["content"])
        accepted = confidence >= JEV_THRESHOLD
        tqdm.write(f"Jev | amount  {row_id:<8} '{amount}' → {choice}  {confidence:.2f}  " f"{'accepted' if accepted else 'review'}")
        jev_stats["asked"] += 1
        if accepted:
            jev_stats["accepted"] += 1
        if confidence >= JEV_THRESHOLD:
            if choice == "thousands":
                result = big_value
            issues.append({"office": office, "id": row_id,
                           "problem": f"Ambiguous amount '{amount}': Jev chose {choice} " f"(confidence {confidence:.2f}), read as {result}"})
        else:
            issues.append({"office": office, "id": row_id,
                "problem": f"Ambiguous amount '{amount}': could be {big_value} or {result}. "
                            f"Jev was not confident ({confidence:.2f}). Amount left empty. Please check."})
            return None
    return result

def parse_phone(row):
    value = row["phone"]
    office = row["office"]
    row_id = row["id"]

    if not isinstance(value, str):
        if row["source_file"] not in files_without_phone:
            issues.append({"office": office, "id": row_id, "problem": "Phone is missing"})
        return None

    try:
        number = phonenumbers.parse(value, office_regions[office])
    except phonenumbers.NumberParseException:
        issues.append({"office": office, "id": row_id, "problem": f"Could not read phone '{value}'"})
        return None

    if not phonenumbers.is_valid_number(number):
        issues.append({"office": office, "id": row_id, "problem": f"Phone '{value}' is not a valid number"})
        return None

    region = phonenumbers.region_code_for_number(number)
    if region != office_regions[office]:
        issues.append({"office": office, "id": row_id, "problem": f"Foreign phone number '{value}' ({region}, office {office})"})

    return phonenumbers.format_number(number, phonenumbers.PhoneNumberFormat.INTERNATIONAL)


def is_japanese(text):
    for c in text:
        char_name = unicodedata.name(c, "")
        if "CJK" in char_name or "HIRAGANA" in char_name or "KATAKANA" in char_name:
            return True
    return False

def split_japanese_name(name):
    parts = name.split(" ")
    if len(parts) == 2:
        return parts[0], parts[1]
    return parts[0], None

def resolve_japanese_name(name, office, row_id):
    if len(name) == 1:
        issues.append({"office": office, "id": row_id,
                       "problem": f"Name '{name}' has only one character (family name only?)"})
        return name, None

    choice, confidence = ask_jev_name(name)

    accepted = confidence >= JEV_THRESHOLD
    tqdm.write(f"Jev | name    {row_id:<8} '{name}' → {choice}  {confidence:.2f}  " f"{'accepted' if accepted else 'review'}")
    jev_stats["asked"] += 1
    if accepted:
        jev_stats["accepted"] += 1

    if confidence < JEV_THRESHOLD:
        issues.append({"office": office, "id": row_id,
                       "problem": f"Could not split name '{name}'. "
                                f"Jev was not confident ({confidence:.2f}). Please check."})
        return name, None

    if choice == "family_only":
        issues.append({"office": office, "id": row_id, "problem": f"Name '{name}': Jev judged it as a family name only "
                                f"(confidence {confidence:.2f})"})
        return name, None

    i = int(choice.split("_")[1])
    family = name[:i]
    given = name[i:]
    issues.append({"office": office, "id": row_id, "problem": f"Name '{name}': Jev split it as {family} / {given} "
                              f"(confidence {confidence:.2f})"})
    return family, given

def split_latin_name(name):
    if name.islower():
        name = name.title()
    parts = name.split(" ")

    titles = ["Mr.", "Mrs.", "Ms.", "Dr.", "Prof.", "M.", "Mme", "Mlle"]
    suffixes = ["Jr.", "Sr.", "II", "III"]

    if parts[0] in titles:
        parts = parts[1:]
    if parts[-1] in suffixes:
        parts = parts[:-1]
        
    if len(parts) == 1:
        return parts[0], None

    upper_words = []
    other_words = []

    for word in parts:
        if word.isupper() and len(word) > 1:
            upper_words.append(word)
        else:
            other_words.append(word)

    if len(upper_words) > 0 and len(other_words) > 0:
        family = " ".join(upper_words).title()
        given = " ".join(other_words)
        return family, given
    
    particles = ["de", "du", "des", "la", "le", "van", "von", "der", "di", "da"]
    for i in range(1, len(parts)):
        if parts[i] in particles:
            family = " ".join(parts[i:])
            given = " ".join(parts[:i])
            return family, given

    family = parts[-1]
    given = " ".join(parts[:-1])
    return family, given

def split_name(row):
    name = row["name"]
    office = row["office"]
    row_id = row["id"]

    if not isinstance(name, str):
        issues.append({"office": office, "id": row_id, "problem": "Name is missing"})
        return pd.Series({"family_name": None, "given_name": None})
    if is_japanese(name):
        family, given = split_japanese_name(name)
        if given is None:
            family, given = resolve_japanese_name(name, office, row_id)
            return pd.Series({"family_name": family, "given_name": given})
    else:
        family, given = split_latin_name(name)
    if given is None:
        issues.append({"office": office, "id": row_id, "problem": f"Could not split name '{name}'"})
    return pd.Series({"family_name": family, "given_name": given})

def write_summary_table(writer, table, title, title_row):
    table.reset_index().to_excel(writer, sheet_name="Summary", startrow=title_row, index=False)

    sheet = writer.sheets["Summary"]
    sheet[f"A{title_row}"] = title
    sheet[f"A{title_row}"].font = Font(bold=True, size=13, color="1F4E78")

    header_row = title_row + 1
    for cell in sheet[header_row]:
        if cell.value is not None:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", start_color="1F4E78")
    return header_row

jev_stats = {"asked": 0, "accepted": 0}
issues = []
tables = []
files_without_phone = []

for path in sorted(pat("input").glob("*.xlsx")):
    if path.name.startswith("~$"):
        continue
    office=find_office(path.stem)
    if office is None:
        issues.append({"office": "(unknown)", "id": f"(whole file: {path.name})", "problem": "Unkown office in file name. File skipped"})
        continue

    table = pd.read_excel(path)

    for col in table.columns:
        if col not in column_names:
            issues.append({"office": office, "id": f"(whole file: {path.name})", "problem": f"Unknown column '{col}'"})

    table = table.rename(columns=column_names)
    table["office"] = office
    table["source_file"] = path.name

    for col in required_columns:
        if col not in table.columns:
            issues.append({"office": office, "id": f"(whole file: {path.name})", "problem": f"Missing column '{col}'"})

    if "phone" not in table.columns:
        files_without_phone.append(path.name)

    tables.append(table)

all_data = pd.concat(tables, ignore_index=True)
for col in ["id", "date", "name", "company", "phone", "amount"]:
    all_data[col] = all_data[col].map(clean_text)

duplicates = all_data[all_data.drop(columns=["source_file"]).duplicated()]

for row_id, office in zip(duplicates["id"], duplicates["office"]):
    issues.append({"office": office, "id": row_id, "problem": "Duplicate row removed"})
all_data = all_data[~all_data.drop(columns=["source_file"]).duplicated()].reset_index(drop=True)

duplicate_ids = all_data[all_data["id"].duplicated(keep=False)]
for row_id, office in zip(duplicate_ids["id"], duplicate_ids["office"]):
    issues.append({"office": office, "id": row_id, "problem": "Same ID used for different inquiries"})

file_date_orders = find_file_date_orders(all_data)

tqdm.pandas(desc="Dates   ")
all_data["date_clean"] = all_data.progress_apply(parse_date, axis=1)
tqdm.pandas(desc="Currency")
all_data["currency"] = all_data.progress_apply(parse_currency, axis=1)
tqdm.pandas(desc="Amounts ")
all_data["amount_clean"] = all_data.progress_apply(parse_amount, axis=1)
tqdm.pandas(desc="Phones  ")
all_data["phone_clean"] = all_data.progress_apply(parse_phone, axis=1)
tqdm.pandas(desc="Names   ")
all_data[["family_name", "given_name"]] = all_data.progress_apply(split_name, axis=1)

issues_table = pd.DataFrame(issues)

def issue_level(problem):
    if "Duplicate row removed" in problem:
        return "Info"
    if "Foreign phone" in problem:
        return "Info"
    if "not confident" in problem:
        return "Error"
    if "Jev" in problem:
        return "AI"
    if "confirmed by the other dates" in problem:
        return "Info"
    if "Ambiguous" in problem or "Year missing" in problem or "does not follow" in problem:
        return "Check"
    return "Error"

def issue_type(problem):
    if "confirmed by the other dates" in problem:
        return "date order confirmed from the same file"
    if "Ambiguous date" in problem:
        return "ambiguous date"
    if "Phone is missing" in problem:
        return "phone missing"
    if "Duplicate row removed" in problem:
        return "duplicate row"
    if "Foreign phone" in problem:
        return "foreign phone number"
    return None

issues_table["level"] = issues_table["problem"].map(issue_level)

issues_table["type"] = issues_table["problem"].map(issue_type)

flagged_rows = set()
ai_rows = set()
for office, row_id, level in zip(issues_table["office"], issues_table["id"], issues_table["level"]):
    if level in ["Error", "Check"]:
        flagged_rows.add((office, row_id))
    elif level == "AI":
        ai_rows.add((office, row_id))


def row_status(row):
    key = (row["office"], row["id"])
    if key in flagged_rows:
        return "Needs review"
    if key in ai_rows:
        return "AI decided"
    return "Clean"


all_data["status"] = all_data.apply(row_status, axis=1)

status_summary = all_data["status"].value_counts().reindex(
    ["Clean", "AI decided", "Needs review"], fill_value=0).to_frame("rows")
status_summary.index.name = "status"

grouped_rows = []
for (office, level, issue_name), group in issues_table.dropna(subset=["type"]).groupby(["office", "level", "type"]):
    examples = ", ".join(list(group["id"])[:3])
    if len(group) > 3:
        examples = examples + ", ..."
    grouped_rows.append({"office": office, "id": "(grouped)", "problem": f"{len(group)} rows: {issue_name} ({examples})", "level": level})

single_rows = issues_table[issues_table["type"].isna()]
issues_summary = pd.concat([single_rows.drop(columns=["type"]), pd.DataFrame(grouped_rows)], ignore_index=True)
issues_summary = issues_summary.sort_values(["office", "level", "problem"], ignore_index=True)

issues_detail = issues_table.drop(columns = ["type"])

office_summary = all_data.groupby("office").agg(
    inquiries=("id", "count"),
)

currency_summary = all_data.groupby("currency").agg(
    inquiries=("amount_clean", "count"),
    total_amount=("amount_clean", "sum"),
)

pending_amounts = all_data[all_data["amount_clean"].isna() & all_data["amount"].notna()]
pending_by_currency = pending_amounts.groupby("currency").size()
currency_summary["pending"] = pending_by_currency.reindex(currency_summary.index, fill_value=0)

issue_summary = pd.crosstab(issues_table["office"], issues_table["level"])

with pd.ExcelWriter("output/master.xlsx", datetime_format="YYYY-MM-DD") as writer:
    office_header = write_summary_table(writer, office_summary, "Inquiries by office", 1)

    next_row = office_header + len(office_summary) + 2
    currency_header = write_summary_table(
        writer, currency_summary,
        "Amount by currency", next_row)

    next_row = currency_header + len(currency_summary) + 2
    issue_header = write_summary_table(writer, issue_summary, "Issues by office and level", next_row)

    summary_sheet = writer.sheets["Summary"]

    for col in ["A", "B", "C", "D", "E"]:
        summary_sheet.column_dimensions[col].width = 14

    for row in range(currency_header + 1, currency_header + 1 + len(currency_summary)):
        summary_sheet[f"C{row}"].number_format = "#,##0.00"
        pending_area = f"D{currency_header + 1}:D{currency_header + len(currency_summary)}"
        pending_fill = PatternFill("solid", start_color="FFE699", end_color="FFE699")
    summary_sheet.conditional_formatting.add(pending_area, CellIsRule(operator="greaterThan", formula=["0"], fill=pending_fill))

    for cell in summary_sheet[issue_header]:
        if cell.value in level_colors:
            cell.fill = PatternFill("solid", start_color=level_colors[cell.value])
            cell.font = Font(bold=True)

        next_row = issue_header + len(issue_summary) + 2
        status_header = write_summary_table(writer, status_summary, "Rows by status", next_row)

        chart = PieChart()
        chart.title = "Rows by status"
        chart.add_data(Reference(summary_sheet, min_col=2, min_row=status_header,
                                max_row=status_header + len(status_summary)), titles_from_data=True)
        chart.set_categories(Reference(summary_sheet, min_col=1, min_row=status_header + 1, 
                                       max_row=status_header + len(status_summary)))
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.dataLabels.showVal = False
        chart.dataLabels.showCatName = False
        chart.dataLabels.showSerName = False

        for i, color in enumerate(["C6E0B4", "BDD7EE", "FFE699"]):
            point = DataPoint(idx=i)
            point.graphicalProperties.solidFill = color
            chart.series[0].data_points.append(point)

    summary_sheet.add_chart(chart, "F2")

    all_data.to_excel(writer, sheet_name="Data", index=False)

    issues_summary.to_excel(writer, sheet_name="Issues", index=False)
    issues_detail.to_excel(writer, sheet_name="Issues detail", index=False)

    for sheet_name in ["Data", "Issues", "Issues detail"]:
        sheet = writer.sheets[sheet_name]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for column in sheet.columns:
            max_length = 0
            for cell in column:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            sheet.column_dimensions[column[0].column_letter].width = min(max_length + 2, 50)

            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", start_color="1F4E78")

        for sheet_name in ["Issues", "Issues detail"]:
            sheet = writer.sheets[sheet_name]
            area = f"A2:D{sheet.max_row}"
            for level, color in level_colors.items():
                fill = PatternFill("solid", start_color=color, end_color=color)
                rule = FormulaRule(formula=[f'$D2="{level}"'], fill=fill)
                sheet.conditional_formatting.add(area, rule)

print(f"Jev: {jev_stats['asked']} questions, {jev_stats['accepted']} accepted, " f"{jev_stats['asked'] - jev_stats['accepted']} sent to review")       
print("Saved: output/master.xlsx")