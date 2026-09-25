import pandas as pd
import unicodedata
import dateparser
from datetime import datetime
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import FormulaRule
from dotenv import load_dotenv
from typesafe_sdk import TypeSafeClient, Choice, TypeSafeError
import phonenumbers

pd.set_option("display.max_columns", None)
pd.set_option("display.width", 200)

tokyo_columns = {
    "問い合わせID": "id",
    "受付日": "date",
    "氏名": "name",
    "会社名": "company",
    "電話番号": "phone",
    "金額": "amount",
    "問い合わせ内容": "content",
}

paris_columns = {
    "N° demande": "id",
    "Date de réception": "date",
    "Nom complet": "name",
    "Société": "company",
    "Montant": "amount",
    "Message": "content",
}

newyork_columns = {
    "Inquiry ID": "id",
    "Date Received": "date",
    "Full Name": "name",
    "Company": "company",
    "Phone": "phone",
    "Amount": "amount",
    "Message": "content",
}

required_columns = ["id", "date", "name", "company", "phone", "amount", "content"]

date_orders = {
    "JP": "YMD",
    "FR": "DMY",
    "US": "MDY",
}

default_currencies = {
    "JP": "JPY",
    "FR": "EUR",
    "US": "USD",
}

decimal_marks = {
    "JP": ".",
    "FR": ",",
    "US": ".",
}

office_regions = {
    "Tokyo": "JP",
    "Paris": "FR",
    "New York": "US",
}

level_colors = {
    "Error": "F8CBAD",
    "Check": "FFE699",
    "Info": "D9D9D9",
    "AI": "BDD7EE"
}

JEV_THRESHOLD = 0.7

load_dotenv()
jev = TypeSafeClient()


def clean_text(value):
    if isinstance(value, str):
        return unicodedata.normalize("NFKC", value).strip()
    return value


def parse_date(row):
    value = row["date"]
    office = row["office"]
    row_id = row["id"]

    if isinstance(value, datetime):
        return value
    if not isinstance(value, str):
        issues.append({"office": office, "id": row_id,
                       "problem": "Date is missing"})
        return None

    if value.startswith("R"):
        year, month, day = value[1:].split(".")
        return datetime(2018 + int(year), int(month), int(day))

    if value[:4].isdigit():
        order = "YMD"
    else:
        order = date_orders[office_regions[office]]

    result = dateparser.parse(
        value,
        languages=["ja", "fr", "en"],
        settings={"DATE_ORDER": order},
    )

    if result is None:
        issues.append({"office": office, "id": row_id,
                       "problem": f"Could not read date '{value}'"})
        return None

    if str(result.year) not in value:
        issues.append({"office": office, "id": row_id,
                       "problem": f"Year missing or abbreviated in '{value}' (assumed {result.year})"})

    parts = value.split("/")
    if len(parts) == 3 and order in ["MDY", "DMY"]:
        first = int(parts[0])
        second = int(parts[1])
        if order == "MDY":
            expected = result.month
        else:
            expected = result.day
        if first != expected:
            issues.append({"office": office, "id": row_id,
                           "problem": f"Date '{value}' does not follow the {office} format ({order})"})

        if first <= 12 and second <= 12 and first != second:
            if order == "MDY":
                other_reading = datetime(result.year, second, first)
            else:
                other_reading = datetime(result.year, first, second)
            if other_reading <= datetime.now():
                issues.append({"office": office, "id": row_id, "problem": f"Ambiguous date '{value}': read as {result:%Y-%m-%d} " f"({office} format), could also be {other_reading:%Y-%m-%d}"})

    if result > datetime.now():
        issues.append({"office": office, "id": row_id,
                       "problem": f"Date '{value}' is in the future (read as {result:%Y-%m-%d})"})

    return result


def parse_currency(row):
    office = row["office"]
    amount = row["amount"]

    if pd.isna(amount):
        return None
    if not isinstance(amount, str):
        return default_currencies[office_regions[office]]

    if "¥" in amount or "円" in amount:
        return "JPY"
    elif "€" in amount:
        return "EUR"
    elif "$" in amount or "USD" in amount:
        return "USD"
    else:
        return default_currencies[office_regions[office]]

def ask_jev_value(office, big_value, small_value, content):
    state = (f"Office: {office}. A number in this message could mean either "
             f"{big_value} or {small_value}. Message: {content}")
    try:
        response = jev.system_one(
            state=state,
            questions={
                "value": Choice(
                    instructions="Which value does the customer mean?",
                    criteria={
                        "thousands": f"{big_value} (a large number)",
                        "decimal": f"{small_value} (a small number with decimals)",
                    },
                ),
            },
        )
    except TypeSafeError as error:
        print(f"Jev error: {error}")
        return None, 0.0

    answer = response.answers["value"]
    return answer.choice, answer.confidence

def name_candidates(name):
    candidates = {}
    for i in range(1, len(name)):
        candidates[f"split_{i}"] = f"Family name: {name[:i]} / Given name: {name[i:]}"
    candidates["family_only"] = f"Family name only: {name} (the given name is missing)"
    return candidates

def ask_jev_name(name):
    candidates = name_candidates(name)
    state = f"A Japanese person's name was written without a space between family name and given name: {name}"
    try:
        response = jev.system_one(
            state=state,
            questions={
                "split": Choice(
                    instructions="How should this name be divided into family name and given name?",
                    criteria=candidates,
                ),
            },
        )
    except TypeSafeError as error:
        print(f"Jev error: {error}")
        return None, 0.0

    answer = response.answers["split"]
    return answer.choice, answer.confidence


def parse_amount(row):
    office = row["office"]
    amount = row["amount"]
    row_id = row["id"]

    if pd.isna(amount):
        issues.append({"office": office, "id": row_id, "problem": "Amount is missing"})
        return None
    if not isinstance(amount, str):
        return float(amount)

    number_text = ""
    for c in amount:
        if c in "0123456789,.-":
            number_text = number_text + c

    if "," in number_text and "." in number_text:
        if number_text.rfind(",") > number_text.rfind("."):
            decimal = ","
        else:
            decimal = "."
    else:
        decimal = decimal_marks[office_regions[office]]

    parts = number_text.split(decimal)
    ambiguous = len(parts) == 2 and len(parts[1]) == 3

    if decimal == ",":
        number_text = number_text.replace(".", "")
        number_text = number_text.replace(",", ".")
    else:
        number_text = number_text.replace(",", "")


    try:
        result = float(number_text)
    except ValueError:
        issues.append({"office": office, "id": row_id, "problem": f"Could not read amount '{amount}'"})
        return None

    if result < 0:
        issues.append({"office": office, "id": row_id, "problem": f"Amount is negative ({amount})"})
    if ambiguous:
        print("JEV is thinking...")
        big_value = int(round(result * 1000))
        choice, confidence = ask_jev_value(office,big_value, result, row["content"])
        if confidence >= JEV_THRESHOLD:
            if choice == "thousands":
                result = big_value
            issues.append({"office": office, "id": row_id,
                           "problem": f"Ambiguous amount '{amount}': Jev chose {choice} "
                                      f"(confidence {confidence:.2f}), read as {result}"})
        else:
            issues.append({"office": office, "id": row_id,
                           "problem": f"Ambiguous amount '{amount}' (read as {result}). "
                                      f"JEV was not confident! ({confidence:.2f}). Please check."})

    return result

def parse_phone(row):
    value = row["phone"]
    office = row["office"]
    row_id = row["id"]

    if not isinstance(value, str):
        if office not in office_without_phone:
            issues.append({"office": office, "id": row_id, "problem": "Phone is missing"})
        return None

    try:
        number = phonenumbers.parse(value, office_regions[office])
    except phonenumbers.NumberParseException:
        issues.append({"office": office, "id": row_id, "problem": f"Could not read phone '{value}'"})
        return None

    if not phonenumbers.is_valid_number(number):
        issues.append({"office": office, "id": row_id, "problem": f"Phone '{value}' is not a valid number"})
        return None

    region = phonenumbers.region_code_for_number(number)
    if region != office_regions[office]:
        issues.append({"office": office, "id": row_id, "problem": f"Foreign phone number '{value}' ({region}, office {office})"})

    return phonenumbers.format_number(number, phonenumbers.PhoneNumberFormat.INTERNATIONAL)


def is_japanese(text):
    for c in text:
        char_name = unicodedata.name(c, "")
        if "CJK" in char_name or "HIRAGANA" in char_name or "KATAKANA" in char_name:
            return True
    return False

def split_japanese_name(name):
    parts = name.split(" ")
    if len(parts) == 2:
        return parts[0], parts[1]
    return parts[0], None

def resolve_japanese_name(name, office, row_id):
    if len(name) == 1:
        issues.append({"office": office, "id": row_id,
                       "problem": f"Name '{name}' has only one character (family name only?)"})
        return name, None

    choice, confidence = ask_jev_name(name)

    if confidence < JEV_THRESHOLD:
        issues.append({"office": office, "id": row_id,
                       "problem": f"Could not split name '{name}'. "
                                f"Jev was not confident ({confidence:.2f}). Please check."})
        return name, None

    if choice == "family_only":
        issues.append({"office": office, "id": row_id, "problem": f"Name '{name}': Jev judged it as a family name only "
                                f"(confidence {confidence:.2f})"})
        return name, None

    i = int(choice.split("_")[1])
    family = name[:i]
    given = name[i:]
    issues.append({"office": office, "id": row_id, "problem": f"Name '{name}': Jev split it as {family} / {given} "
                              f"(confidence {confidence:.2f})"})
    return family, given

def split_latin_name(name):
    if name.islower():
        name = name.title()
    parts = name.split(" ")

    titles = ["Mr.", "Mrs.", "Ms.", "Dr.", "Prof.", "M.", "Mme", "Mlle"]
    suffixes = ["Jr.", "Sr.", "II", "III"]

    if parts[0] in titles:
        parts = parts[1:]
    if parts[-1] in suffixes:
        parts = parts[:-1]
        
    if len(parts) == 1:
        return parts[0], None

    upper_words = []
    other_words = []

    for word in parts:
        if word.isupper() and len(word) > 1:
            upper_words.append(word)
        else:
            other_words.append(word)

    if len(upper_words) > 0 and len(other_words) > 0:
        family = " ".join(upper_words).title()
        given = " ".join(other_words)
        return family, given
    
    particles = ["de", "du", "des", "la", "le", "van", "von", "der", "di", "da"]
    for i in range(1, len(parts)):
        if parts[i] in particles:
            family = " ".join(parts[i:])
            given = " ".join(parts[:i])
            return family, given

    family = parts[-1]
    given = " ".join(parts[:-1])
    return family, given

def split_name(row):
    name = row["name"]
    office = row["office"]
    row_id = row["id"]

    if not isinstance(name, str):
        issues.append({"office": office, "id": row_id, "problem": "Name is missing"})
        return pd.Series({"family_name": None, "given_name": None})
    if is_japanese(name):
        family, given = split_japanese_name(name)
        if given is None:
            family, given = resolve_japanese_name(name, office, row_id)
            return pd.Series({"family_name": family, "given_name": given})
    else:
        family, given = split_latin_name(name)
    if given is None:
        issues.append({"office": office, "id": row_id, "problem": f"Could not split name '{name}'"})
    return pd.Series({"family_name": family, "given_name": given})

tokyo = pd.read_excel("input/inquiries_tokyo.xlsx")
paris = pd.read_excel("input/inquiries_paris.xlsx")
newyork = pd.read_excel("input/inquiries_newyork.xlsx")

tokyo = tokyo.rename(columns=tokyo_columns)
tokyo["office"] = "Tokyo"

paris = paris.rename(columns=paris_columns)
paris["office"] = "Paris"

newyork = newyork.rename(columns=newyork_columns)
newyork["office"] = "New York"

tables = {
    "Tokyo": tokyo,
    "Paris": paris,
    "New York": newyork,
}

office_without_phone = []
for office, table in tables.items():
    if "phone" not in table.columns:
        office_without_phone.append(office)

issues = []

for office, table in tables.items():
    for col in required_columns:
        if col not in table.columns:
            issues.append({
                "office": office,
                "id": "(whole file)",
                "problem": f"Missing column '{col}'"
            })

all_data = pd.concat([tokyo, paris, newyork], ignore_index=True)

for col in ["id", "date", "name", "company", "phone", "amount"]:
    all_data[col] = all_data[col].map(clean_text)

duplicates = all_data[all_data.duplicated()]
for row_id, office in zip(duplicates["id"], duplicates["office"]):
    issues.append({"office": office, "id": row_id, "problem": "Duplicate row removed"})
all_data = all_data.drop_duplicates(ignore_index=True)

duplicate_ids = all_data[all_data["id"].duplicated(keep=False)]
for row_id, office in zip(duplicate_ids["id"], duplicate_ids["office"]):
    issues.append({"office": office, "id": row_id, "problem": "Same ID used for different inquiries"})

all_data["date_clean"] = all_data.apply(parse_date, axis=1)
all_data["currency"] = all_data.apply(parse_currency, axis=1)
all_data["amount_clean"] = all_data.apply(parse_amount, axis=1)
all_data["phone_clean"] = all_data.apply(parse_phone, axis=1)
all_data[["family_name", "given_name"]] = all_data.apply(split_name, axis=1)

issues_table = pd.DataFrame(issues)

def issue_level(problem):
    if "Duplicate row removed" in problem:
        return "Info"
    if "Foreign phone" in problem:
        return "Info"
    if "not confident" in problem:
        return "Error"
    if "Jev" in problem:
        return "AI"
    if "Ambiguous" in problem or "Year missing" in problem or "does not follow" in problem or "negative" in problem:
        return "Check"
    return "Error"

issues_table["level"] = issues_table["problem"].map(issue_level)

office_summary = all_data.groupby("office").agg(
    inquiries=("id", "count"),
)

currency_summary = all_data.groupby("currency").agg(
    inquiries=("id", "count"),
    total_amount=("amount_clean", "sum"),
)

issue_summary = pd.crosstab(issues_table["office"], issues_table["level"])

with pd.ExcelWriter("output/master.xlsx", datetime_format="YYYY-MM-DD") as writer:
    office_summary.reset_index().to_excel(writer, sheet_name="Summary", startrow=1, index=False)
    currency_summary.reset_index().to_excel(writer, sheet_name="Summary", startrow=7, index=False)
    issue_summary.reset_index().to_excel(writer, sheet_name="Summary", startrow=13, index=False)

    summary_sheet = writer.sheets["Summary"]
    summary_sheet["A1"] = "Inquiries by office"
    summary_sheet["A7"] = "Amount by currency (amounts not yet checked, including negative values)"
    summary_sheet["A13"] = "Issues by office and level"

    summary_sheet.column_dimensions["A"].width = 14
    for col in ["B", "C", "D", "E"]:
        summary_sheet.column_dimensions[col].width = 14

    title_font = Font(bold=True, size=13, color="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="1F4E78")

    for row in [1, 7, 13]:
        summary_sheet[f"A{row}"].font = title_font

    for row in [2, 8, 14]:
        for cell in summary_sheet[row]:
            if cell.value is not None:
                cell.font = header_font
                cell.fill = header_fill

    for cell in summary_sheet[14]:
        if cell.value in level_colors:
            cell.fill = PatternFill("solid", start_color=level_colors[cell.value])
            cell.font = Font(bold=True)
    
    for row in range(9, 12):
        summary_sheet[f"C{row}"].number_format = "#,##0.00"

    all_data.to_excel(writer, sheet_name="Data", index=False)
    issues_table.to_excel(writer, sheet_name="Issues", index=False)

    for sheet_name in ["Data", "Issues"]:
        sheet = writer.sheets[sheet_name]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for column in sheet.columns:
            max_length = 0
            for cell in column:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            sheet.column_dimensions[column[0].column_letter].width = min(max_length + 2, 50)

            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", start_color="1F4E78")

    issues_sheet = writer.sheets["Issues"]
    area = f"A2:D{issues_sheet.max_row}"

    for level, color in level_colors.items():
        fill = PatternFill("solid", start_color=color, end_color=color)
        rule = FormulaRule(formula=[f'$D2="{level}"'], fill=fill)
        issues_sheet.conditional_formatting.add(area, rule)
        
print("Saved: output/master.xlsx")