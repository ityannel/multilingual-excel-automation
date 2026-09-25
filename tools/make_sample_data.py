"""サンプルデータ生成スクリプト

東京・パリ・ニューヨークの3拠点から届いた問い合わせExcel（input/）と、
答え合わせ用の正解表（answer_key/answer_key.xlsx）を作成する。

実行:  .venv\\Scripts\\python.exe tools\\make_sample_data.py
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
INPUT_DIR = ROOT / "input"
KEY_DIR = ROOT / "answer_key"

# 正解（truth）の各項目:
#   date: 正しい受付日（YYYY-MM-DD、決められない場合は None）
#   lang: ja / fr / en / mixed / unknown
#   family, given: 姓・名
#   company: 正規化後の会社名
#   amount, currency: 数値と通貨
#   category: 請求 / 配送 / 技術 / 解約 / 見積もり / その他
#   issues: このデータに意図的に入れた問題
#   note: 補足


def T(date, lang, family, given, company, amount, currency, category, issues=(), note=""):
    return dict(date=date, lang=lang, family=family, given=given, company=company,
                amount=amount, currency=currency, category=category,
                issues=list(issues), note=note)


SAKURA = "株式会社サクラ商事"
MIDORI = "株式会社ミドリ"
AOBA = "株式会社アオバ"
SOCGEN = "Société Générale"
ACME = "Acme Corp"

MSG_DUPONT = ("Bonjour, je souhaite résilier notre contrat à la fin du mois prochain. "
              "Merci de me confirmer la procédure.")
MSG_SMITH = "Could you send us a revised quotation for 200 units? We need it by Friday."

# ---------------------------------------------------------------- 東京（日本語の見出し）
TOKYO_HEADERS = ["問い合わせID", "受付日", "氏名", "会社名", "電話番号", "金額", "問い合わせ内容"]
TOKYO = [
    ("T-001", "2026/07/03", "山田 太郎", "株式会社サクラ商事", "03-1234-5678", "¥12,000",
     "7月分の請求書の金額が見積もりと違っています。確認をお願いします。",
     T("2026-07-03", "ja", "山田", "太郎", SAKURA, 12000, "JPY", "請求")),
    ("T-002", "2026年7月8日", "佐藤花子", "株式会社ミドリ", "06-2345-6789", "8500",
     "注文した商品がまだ届きません。配送状況を教えてください。",
     T("2026-07-08", "ja", "佐藤", "花子", MIDORI, 8500, "JPY", "配送",
       ["氏名に区切りなし"], "姓と名の間に空白がない")),
    ("T-003", "２０２６年７月１５日", "  鈴木　一郎 ", "ｻｸﾗ商事", "０３－１２３４－５６７８", "１５，０００円",
     "ログインするとエラー画面が表示されます。パスワードを再設定しても直りません。",
     T("2026-07-15", "ja", "鈴木", "一郎", SAKURA, 15000, "JPY", "技術",
       ["全角文字", "余分な空白", "半角カナ", "会社名の表記ゆれ"], "全角スペースを含む")),
    ("T-004", "R8.7.30", "Marie Dupont", "Société Générale", "+33 1 23 45 67 89", "¥45,000",
     MSG_DUPONT,
     T("2026-07-30", "fr", "Dupont", "Marie", SOCGEN, 45000, "JPY", "解約",
       ["和暦"], "令和8年＝2026年。P-002と同じ問い合わせ")),
    ("T-005", "2026-08-04", "John Smith", "Acme Corp.", None, "23000",
     MSG_SMITH,
     T("2026-08-04", "en", "Smith", "John", ACME, 23000, "JPY", "見積もり",
       ["電話番号欠損", "会社名の表記ゆれ"], "N-004と同じ問い合わせの可能性")),
    ("T-006", "2026/08/12", "高橋 美咲", "株式会社サクラ商事 ", "03-1234-5678", "-3000",
     "先月の請求で二重に引き落とされています。返金をお願いします。",
     T("2026-08-12", "ja", "高橋", "美咲", SAKURA, -3000, "JPY", "請求",
       ["金額が負", "余分な空白"])),
    ("T-006", "2026/08/12", "高橋 美咲", "株式会社サクラ商事 ", "03-1234-5678", "-3000",
     "先月の請求で二重に引き落とされています。返金をお願いします。",
     T("2026-08-12", "ja", "高橋", "美咲", SAKURA, -3000, "JPY", "請求",
       ["重複行"], "直前の行と完全に同じ")),
    ("T-008", "8/19", "伊藤 健", "株式会社アオバ", "045-111-2222", "5,000",
     "来週の会議の前に、新しいプランについて説明していただけますか。",
     T("2026-08-19", "ja", "伊藤", "健", AOBA, 5000, "JPY", "その他",
       ["日付の年なし"], "前後の行から2026年と推定")),
    ("T-009", "2026年8月25日", "田中", "株式会社ミドリ", None, "7000",
     "配達日を9月3日に変更できますか。",
     T("2026-08-25", "ja", "田中", None, MIDORI, 7000, "JPY", "配送",
       ["名の欠損", "電話番号欠損"], "本文中の「9月3日」は受付日ではない")),
    ("T-010", "2026/09/01", "YAMADA Taro", "Sakura Shoji Co., Ltd.", "03-1234-5678", "¥12,000",
     "I have not received the invoice for August yet. Please resend it.",
     T("2026-09-01", "en", "Yamada", "Taro", SAKURA, 12000, "JPY", "請求",
       ["氏名の順序", "会社名の表記ゆれ"], "ローマ字・姓が先（大文字）。会社と電話番号からT-001の山田 太郎と同一人物の可能性")),
    ("T-011", "2026/09/02", "渡辺 直樹", "株式会社アオバ", "045-111-2222", None,
     "アプリが頻繁に落ちます。Android 15で使用しています。",
     T("2026-09-02", "ja", "渡辺", "直樹", AOBA, None, None, "技術", ["金額欠損"])),
    ("T-012", "2026/09/07", "小林 さくら", None, "090-1234-5678", "2000",
     "Merci pour votre aide. 問題は解決しました。",
     T("2026-09-07", "mixed", "小林", "さくら", None, 2000, "JPY", "その他",
       ["会社名欠損", "言語混在"], "フランス語＋日本語")),
    ("Ｔ－０１３", "2026/9/10", "中村 優", "株式会社ミドリ", "06-2345-6789", "9800",
     "解約の手続き方法を教えてください。今月末で終了したいです。",
     T("2026-09-10", "ja", "中村", "優", MIDORI, 9800, "JPY", "解約", ["ID全角"])),
    ("T-014", datetime(2026, 9, 15), "山本 翔", "株式会社サクラ商事", "03-1234-5678", 12000,
     "見積書の有効期限を延長していただけますか。",
     T("2026-09-15", "ja", "山本", "翔", SAKURA, 12000, "JPY", "見積もり",
       [], "日付・金額がExcelの日付型・数値型で入っている")),
    ("T-015", "2026/09/18", "加藤 恵", "株式会社ミドリ", "06-2345-6789", "4000",
     "ok",
     T("2026-09-18", "unknown", "加藤", "恵", MIDORI, 4000, "JPY", "その他",
       ["言語判定困難（短文）"])),
]

# ---------------------------------------------------------------- パリ（フランス語の見出し、電話番号の列なし）
PARIS_HEADERS = ["N° demande", "Date de réception", "Nom complet", "Société", "Montant", "Message"]
PARIS = [
    ("P-001", "03/07/2026", "Jean-Pierre de la Fontaine", "Société Générale", "1 234,50 €",
     "Bonjour, la facture n° 4521 comporte une erreur sur le montant de la TVA. Pouvez-vous la corriger ?",
     T("2026-07-03", "fr", "de la Fontaine", "Jean-Pierre", SOCGEN, 1234.50, "EUR", "請求",
       ["複合姓"], "日/月/年の順。「de la」は姓の一部")),
    ("P-002", "30/07/2026", "DUPONT Marie", "Societe Generale", None,
     MSG_DUPONT,
     T("2026-07-30", "fr", "Dupont", "Marie", SOCGEN, None, None, "解約",
       ["他ファイルとの重複の可能性", "氏名の順序", "会社名の表記ゆれ", "金額欠損"],
       "T-004と同じ問い合わせ。アクセントなしの会社名")),
    ("P-003", "07/08/2026", "Camille Martin", "Boulangerie Martin & Fils", "2.500,00",
     "Notre commande n'est toujours pas arrivée. Le suivi indique « en transit » depuis dix jours.",
     T("2026-08-07", "fr", "Martin", "Camille", "Boulangerie Martin & Fils", 2500.00, "EUR", "配送",
       ["日付が曖昧"], "パリ拠点なので8月7日（米国式なら7月8日）")),
    ("P-004", "1er août 2026", "Élodie Lefèvre", "Lefèvre SARL", "320,00 €",
     "Impossible de me connecter à mon compte depuis la mise à jour.",
     T("2026-08-01", "fr", "Lefèvre", "Élodie", "Lefèvre SARL", 320.00, "EUR", "技術")),
    ("P-005", "lundi 14 septembre 2026", "Hélène Moreau", "Moreau & Associés", "-120,00 €",
     "Nous avons été facturés deux fois pour le mois d'août.",
     T("2026-09-14", "fr", "Moreau", "Hélène", "Moreau & Associés", -120.00, "EUR", "請求",
       ["金額が負"])),
    ("P-006", "2026-07-21", "Lucas Bernard", "Bernard Transports", "4 500 €",
     "Pourriez-vous nous envoyer un devis pour 50 palettes ?",
     T("2026-07-21", "fr", "Bernard", "Lucas", "Bernard Transports", 4500, "EUR", "見積もり",
       ["日付形式が他と不一致"])),
    ("P-007", "5 septembre 2026", "Mme Sophie Laurent", "Laurent Conseil", "990",
     "Hello, could you tell me when the technician will come? The printer still does not work.",
     T("2026-09-05", "en", "Laurent", "Sophie", "Laurent Conseil", 990, "EUR", "技術",
       ["敬称あり"], "パリ拠点だが本文は英語")),
    ("P-008", "12/08/2026", "Nicolas Petit", "Petit & Cie", "1 100,00 €",
     "Le colis est arrivé endommagé. Je voudrais un remplacement.",
     T("2026-08-12", "fr", "Petit", "Nicolas", "Petit & Cie", 1100.00, "EUR", "配送",
       [], "12月8日と読むと未来の日付になる")),
    ("P-009", "15/09/2026", "Chloé Roux", None, "760 €",
     "Merci de mettre à jour l'adresse de facturation de notre société.",
     T("2026-09-15", "fr", "Roux", "Chloé", None, 760, "EUR", "請求", ["会社名欠損"])),
    ("P-010", "31/09/2026", "Thomas Girard", "Girard Immobilier", "600 €",
     "Pouvez-vous me rappeler demain matin ?",
     T(None, "fr", "Girard", "Thomas", "Girard Immobilier", 600, "EUR", "その他",
       ["無効な日付"], "9月31日は存在しない")),
    ("P-003", "18/09/2026", "Julie Fournier", "Fournier Design", "2 300 €",
     "Je voudrais annuler ma commande n° 7788.",
     T("2026-09-18", "fr", "Fournier", "Julie", "Fournier Design", 2300, "EUR", "解約",
       ["ID重複"], "内容の違う行が同じIDを使っている")),
    ("P-012", "02/09/2026", "Antoine Mercier-Blanc", "Mercier-Blanc SAS", "1 500,00 €",
     "Bonjour, 請求書を日本語でもらえますか ? Merci.",
     T("2026-09-02", "mixed", "Mercier-Blanc", "Antoine", "Mercier-Blanc SAS", 1500.00, "EUR", "請求",
       ["言語混在"], "フランス語＋日本語")),
    ("P-013", " 10/09/2026 ", "yuki tanaka", "Aoba Ltd", "300 €",
     "Bonjour, je travaille à l'antenne de Paris. Le logiciel affiche un message d'erreur au démarrage.",
     T("2026-09-10", "fr", "Tanaka", "Yuki", AOBA, 300, "EUR", "技術",
       ["余分な空白", "小文字の氏名", "会社名の表記ゆれ"],
       "日本人名のローマ字・名が先。姓 Tanaka／名 Yuki")),
    ("P-014", "21/08/2026", "Léa Dubois", "Dubois Événements", "5 000,00 €",
     "Nous aimerions connaître vos tarifs pour une livraison express.",
     T("2026-08-21", "fr", "Dubois", "Léa", "Dubois Événements", 5000.00, "EUR", "見積もり",
       ["分類が曖昧"], "「配送」にも見えるが、料金を尋ねているので見積もり")),
    ("P-015", "28/08/2026", "Hugo Lambert", "Lambert BTP", "abc",
     "Merci.",
     T("2026-08-28", "fr", "Lambert", "Hugo", "Lambert BTP", None, None, "その他",
       ["金額が数値でない", "言語判定困難（短文）"])),
    ("P-016", "16/09/2026", "Kevin Moore", "Moore Consulting", "25,000.12 €",
     "Hi, I am with the Paris branch. Could you confirm the payment schedule for this order?",
     T("2026-09-16", "en", "Moore", "Kevin", "Moore Consulting", 25000.12, "EUR", "請求",
       ["金額が英語式"], "パリ拠点だが英語式（, が桁区切り、. が小数点）。両方あるので後ろの . が小数点と判断できる")),
    ("P-017", "17/09/2026", "Emma Collins", "Collins Travel", "1,500",
     "Please send me the updated invoice for the group booking.",
     T("2026-09-17", "en", "Collins", "Emma", "Collins Travel", 1500, "EUR", "請求",
       ["金額が曖昧"], "区切りが , だけなので 1500（英語式）とも 1.5（フランス式）とも読める。本文が英語なので 1500")),
    ("P-018", "19/09/2026", "Manon Garnier", "Garnier Fleurs", "12,50 €",
     "Bonjour, les frais de port de 12,50 € me semblent incorrects.",
     T("2026-09-19", "fr", "Garnier", "Manon", "Garnier Fleurs", 12.50, "EUR", "請求",
       [], "フランス式の小数。, の後ろが2桁")),
]

# ---------------------------------------------------------------- ニューヨーク（英語の見出し、月/日/年）
NY_HEADERS = ["Inquiry ID", "Date Received", "Full Name", "Company", "Phone", "Amount", "Message"]
NY = [
    ("N-001", "07/08/2026", "Emily Johnson", "Brightside Inc.", "(212) 555-0142", "$1,250.00",
     "We were charged twice for our July subscription. Please refund the duplicate charge.",
     T("2026-07-08", "en", "Johnson", "Emily", "Brightside Inc.", 1250.00, "USD", "請求",
       ["日付が曖昧"], "NY拠点なので7月8日（パリのP-003と同じ文字列だが意味が違う）")),
    ("N-002", "July 14, 2026", "Michael Brown", "Acme Corp", "212-555-0199", "$980",
     "Our shipment was supposed to arrive last Monday but tracking has not updated.",
     T("2026-07-14", "en", "Brown", "Michael", ACME, 980, "USD", "配送")),
    ("N-003", "Aug 3 2026", "Sarah O'Connor", "O'Connor & Partners LLC", "+1 646 555 0110", "2,400",
     "The dashboard keeps timing out when I export reports to Excel.",
     T("2026-08-03", "en", "O'Connor", "Sarah", "O'Connor & Partners LLC", 2400, "USD", "技術",
       ["通貨記号なし"])),
    ("N-004", "08/12/2026", "John Smith", "ACME Corporation", None, "$23,000",
     MSG_SMITH,
     T("2026-08-12", "en", "Smith", "John", ACME, 23000, "USD", "見積もり",
       ["他ファイルとの重複の可能性", "会社名の表記ゆれ", "電話番号欠損"],
       "T-005と同じ本文。受付日と通貨が違うので要確認")),
    ("N-005", "2026-08-20", "Dr. Robert Miller Jr.", "Miller Health", "(718) 555-0123", "$-450.00",
     "Please cancel our plan effective next month.",
     T("2026-08-20", "en", "Miller", "Robert", "Miller Health", -450.00, "USD", "解約",
       ["敬称あり", "金額が負", "日付形式が他と不一致"], "「Dr.」「Jr.」を除く")),
    ("N-006", "8/27/26", "Jessica Davis", "Davis & Co.", "917.555.0177", "$3,100.50",
     "Can you give us a quote for the premium plan with 25 users?",
     T("2026-08-27", "en", "Davis", "Jessica", "Davis & Co.", 3100.50, "USD", "見積もり",
       ["年が2桁"])),
    ("N-007", "September 1st, 2026", "Kenji Watanabe", "Aoba Co., Ltd.", "212-555-0188", "$1,500",
     "ニューヨーク支社の渡辺です。請求書の宛名を変更してください。",
     T("2026-09-01", "ja", "Watanabe", "Kenji", AOBA, 1500, "USD", "請求",
       ["会社名の表記ゆれ"],
       "NY拠点だが本文は日本語。日本人名だが名が先。姓 Watanabe／名 Kenji")),
    ("N-008", "09/03/2026", "Ashley Wilson", "Wilson Logistics", "(347) 555-0166", "$640",
     "The package was delivered to the wrong address.",
     T("2026-09-03", "en", "Wilson", "Ashley", "Wilson Logistics", 640, "USD", "配送")),
    ("N-008", "09/03/2026", "Ashley Wilson  ", "Wilson Logistics", "(347) 555-0166", "$640",
     "The package was delivered to the wrong address.",
     T("2026-09-03", "en", "Wilson", "Ashley", "Wilson Logistics", 640, "USD", "配送",
       ["重複行", "余分な空白"], "空白を除くと直前の行と同じ")),
    ("N-010", "09/10/2026", "Daniel García", "García Imports", "212-555-0101", "$2,000.00",
     "Bonjour, nous avons un problème avec la synchronisation des données depuis hier.",
     T("2026-09-10", "fr", "García", "Daniel", "García Imports", 2000.00, "USD", "技術",
       [], "NY拠点だが本文はフランス語")),
    ("N-011", None, "Olivia Martinez", "Martinez Studio", "646-555-0133", "$870",
     "Could you explain the difference between the Basic and Pro plans?",
     T(None, "en", "Martinez", "Olivia", "Martinez Studio", 870, "USD", "その他", ["日付欠損"])),
    ("N-012", "09/14/2026", "Chris Lee", "Lee & Kim LLC", "212-555-0155", "$1,050",
     "Hi",
     T("2026-09-14", "unknown", "Lee", "Chris", "Lee & Kim LLC", 1050, "USD", "その他",
       ["言語判定困難（短文）"])),
    ("N-013", "09/16/2026", "Matthew Taylor", "taylor media", "347-555-0190", "USD 720",
     "We want to terminate the contract at the end of this quarter.",
     T("2026-09-16", "en", "Taylor", "Matthew", "Taylor Media", 720, "USD", "解約",
       ["会社名が小文字"])),
    ("N-014", "13/09/2026", "Amanda White", "White Consulting", "917-555-0145", "$450",
     "Our invoice shows the wrong company address.",
     T("2026-09-13", "en", "White", "Amanda", "White Consulting", 450, "USD", "請求",
       ["日付形式が他と不一致"], "月/日としては無効。日/月で書かれている")),
    ("N-015", "Sept. 18, 2026", "Ryan Clark", "Clark Engineering", "718-555-0170", "$2,750",
     "Is it possible to get the order delivered before the 25th?",
     T("2026-09-18", "en", "Clark", "Ryan", "Clark Engineering", 2750, "USD", "配送")),
    ("N-016", "09/19/2026", "Pierre Rousseau", "Rousseau Wines", "212-555-0177", "$1.250,00",
     "Bonjour, je travaille au bureau de New York. Pouvez-vous vérifier le montant de cette facture ?",
     T("2026-09-19", "fr", "Rousseau", "Pierre", "Rousseau Wines", 1250.00, "USD", "請求",
       ["金額がフランス式"], "NY拠点だがフランス式（. が桁区切り、, が小数点）。両方あるので後ろの , が小数点と判断できる")),
]

FILES = [
    ("inquiries_tokyo.xlsx", "Tokyo", TOKYO_HEADERS, TOKYO),
    ("inquiries_paris.xlsx", "Paris", PARIS_HEADERS, PARIS),
    ("inquiries_newyork.xlsx", "New York", NY_HEADERS, NY),
]

HEADER_FONT = Font(bold=True)


def write_input(filename, sheet_title, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for row in rows:
        ws.append(list(row[:-1]))  # 最後の要素は正解なので書かない
    for cell in ws[1]:
        cell.font = HEADER_FONT
    for col, width in zip("ABCDEFG", [14, 24, 28, 28, 20, 14, 80]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy/mm/dd"
    wb.save(INPUT_DIR / filename)


def write_answer_key():
    wb = Workbook()
    ws = wb.active
    ws.title = "正解"
    headers = ["ファイル", "Excel行番号", "元のID", "受付日", "言語", "姓", "名", "会社名（正規化）",
               "金額", "通貨", "種類", "意図的に入れた問題", "備考"]
    ws.append(headers)
    issue_counts = {}
    for filename, _, _, rows in FILES:
        for i, row in enumerate(rows):
            t = row[-1]
            for issue in t["issues"]:
                issue_counts[issue] = issue_counts.get(issue, 0) + 1
            ws.append([filename, i + 2, row[0], t["date"], t["lang"], t["family"], t["given"],
                       t["company"], t["amount"], t["currency"], t["category"],
                       "、".join(t["issues"]), t["note"]])

    fill = PatternFill("solid", start_color="D9E1F2")
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = fill
    for col, width in zip("ABCDEFGHIJKLM", [24, 12, 12, 12, 8, 16, 14, 28, 12, 8, 10, 48, 60]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        row[11].alignment = Alignment(wrap_text=True, vertical="top")
        row[12].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("問題の種類")
    ws2.append(["問題の種類", "件数"])
    for issue, count in sorted(issue_counts.items(), key=lambda x: -x[1]):
        ws2.append([issue, count])
    ws2.append([])
    ws2.append(["ファイル単位の問題", ""])
    ws2.append(["inquiries_paris.xlsx に「電話番号」の列がない（必須列の欠損）", 1])
    ws2.append(["3つのファイルで見出しの言語が違う（日本語・フランス語・英語）", 3])
    for cell in ws2[1]:
        cell.font = HEADER_FONT
        cell.fill = fill
    ws2.column_dimensions["A"].width = 60
    ws2.column_dimensions["B"].width = 8

    ws3 = wb.create_sheet("列の対応表")
    ws3.append(["統一後の列名", "東京", "パリ", "ニューヨーク"])
    for unified, *cols in zip(
        ["ID", "受付日", "氏名", "会社名", "電話番号", "金額", "問い合わせ内容"],
        TOKYO_HEADERS,
        PARIS_HEADERS[:4] + ["（なし）"] + PARIS_HEADERS[4:],
        NY_HEADERS,
    ):
        ws3.append([unified, *cols])
    for cell in ws3[1]:
        cell.font = HEADER_FONT
        cell.fill = fill
    for col in "ABCD":
        ws3.column_dimensions[col].width = 22

    wb.save(KEY_DIR / "answer_key.xlsx")
    return issue_counts


def main():
    INPUT_DIR.mkdir(exist_ok=True)
    KEY_DIR.mkdir(exist_ok=True)
    for filename, sheet_title, headers, rows in FILES:
        write_input(filename, sheet_title, headers, rows)
    counts = write_answer_key()
    total = sum(len(rows) for *_, rows in FILES)
    print(f"入力ファイル {len(FILES)} 個、合計 {total} 行を作成しました。")
    print(f"意図的な問題: {len(counts)} 種類、延べ {sum(counts.values())} 件")


if __name__ == "__main__":
    main()
